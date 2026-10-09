from pandas import DataFrame
from pathlib import Path
from typing import Dict, Any, List
import pandas as pd
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import config

class ExcelBuilder:
    def __init__(self, records: List[Dict[str, Any]], output_path: Path):
        self.records = records
        self.output_path = Path(output_path)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
    
    def build(self) -> Path:
        df = pd.DataFrame(self.records)
        with pd.ExcelWriter(self.output_path, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name= "PDF_Form_Data", index=False)
        
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["PDF_Form_Data"]
        self._apply_styling(ws, df)
        wb.save(self.output_path)
        return self.output_path
    
    def _apply_styling(self, ws:Worksheet, df:DataFrame):
        header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        align_center = Alignment(horizontal= "center", vertical= "center", wrap_text=True)
        align_left  = Alignment(vertical="center")
        selected_fill = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
        not_selected_fill = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")
        
        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )

        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 28
        for column_idx in range(1, len(df.columns)+1):
            cell = ws.cell(row=1, column = column_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = thin_border
        
        for row_idx in range(2, len(df)+2):
            ws.row_dimensions[row_idx].height = 20
            for column_idx in range(1,len(df.columns)+1):
                cell = ws.cell(row=row_idx, column = column_idx)
                cell.font = regular_font
                cell.alignment = align_left
                cell.border = thin_border
                column_name = df.columns[column_idx - 1]
                
                value = str(cell.value or "")
                if value == "Selected":
                    cell.fill = selected_fill
                    cell.alignment = align_center
                elif value == "Not Selected":
                    cell.fill = not_selected_fill
                    cell.alignment = align_center
                if column_name == config.SYSTEM_COLUMNS.get("source_file"):
                    pdf_file = config.INPUT_PDF_DIR / value
                    if pdf_file.exists():
                        cell.hyperlink = pdf_file.as_uri()
                        cell.font = link_font
                
                if column_name == config.SYSTEM_COLUMNS.get("signature_link") and value != "N/A":
                    sig_file = Path(value)
                    if sig_file.exists():
                        cell.value = f"View Signature ({sig_file.name})"
                        cell.hyperlink = sig_file.as_uri()
                        cell.font = link_font

        for column in ws.columns:
            max_len = 0
            column_letter = get_column_letter(column[0].column)
            column_header = str(column[0].value or "")
            if column_header == config.SYSTEM_COLUMNS.get("signature_base64"):
                ws.column_dimensions[column_letter].width=25
                continue
            
            for cell in column:
                text = str(cell.value or "")
                if len(text) > max_len:
                    max_len = len(text)
            ws.column_dimensions[column_letter].width = min(max(max_len + 4, 14), 45)